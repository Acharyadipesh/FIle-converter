import os
import uuid
import shutil
import subprocess
import zipfile
from pathlib import Path

from flask import (
    Flask,
    render_template,
    request,
    jsonify,
    send_file
)

from werkzeug.utils import secure_filename

from PIL import Image

import fitz

from pptx import Presentation
from pptx.util import Inches


# ============================================================
# APPLICATION CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"

UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


app = Flask(__name__)

# Maximum upload size = 100 MB
app.config["MAX_CONTENT_LENGTH"] = (
    100 * 1024 * 1024
)


# ============================================================
# ALLOWED FILE TYPES
# ============================================================

ALLOWED = {

    "pdf-to-ppt": {
        ".pdf"
    },

    "ppt-to-pdf": {
        ".ppt",
        ".pptx"
    },

    "jpg-to-png": {
        ".jpg",
        ".jpeg"
    },

    "png-to-jpg": {
        ".png"
    },

    "pdf-to-jpg": {
        ".pdf"
    },

    "jpg-to-pdf": {
        ".jpg",
        ".jpeg"
    },

    "png-to-pdf": {
        ".png"
    },

    "merge-pdf": {
        ".pdf"
    },

    "compress-image": {
        ".jpg",
        ".jpeg",
        ".png",
        ".webp"
    }
}


# ============================================================
# HELPER
# ============================================================

def safe_filename(filename):

    filename = secure_filename(
        filename
    )

    if not filename:
        filename = "file"

    return filename


# ============================================================
# PDF → PPT
# ============================================================

def convert_pdf_to_ppt(
    pdf_file,
    output_file
):

    pdf = fitz.open(
        str(pdf_file)
    )

    if len(pdf) == 0:

        pdf.close()

        raise ValueError(
            "The PDF contains no pages."
        )


    first_page = pdf[0].rect


    presentation = Presentation()


    # Set presentation size according to PDF
    presentation.slide_width = Inches(10)

    presentation.slide_height = Inches(
        10 *
        first_page.height /
        first_page.width
    )


    blank_layout = (
        presentation.slide_layouts[6]
    )


    for page_number, page in enumerate(
        pdf,
        start=1
    ):

        pixmap = page.get_pixmap(
            matrix=fitz.Matrix(
                1.6,
                1.6
            ),
            alpha=False
        )


        temporary_image = (
            OUTPUT_DIR /
            f"temp_{uuid.uuid4().hex}.jpg"
        )


        pixmap.save(
            str(temporary_image)
        )


        slide = (
            presentation
            .slides
            .add_slide(
                blank_layout
            )
        )


        slide.shapes.add_picture(

            str(temporary_image),

            0,

            0,

            width=presentation.slide_width,

            height=presentation.slide_height
        )


        temporary_image.unlink(
            missing_ok=True
        )


    presentation.save(
        str(output_file)
    )


    pdf.close()


# ============================================================
# FIND LIBREOFFICE
# ============================================================

def find_libreoffice():

    # First check PATH
    path_from_system = (
        shutil.which("soffice")
    )

    if path_from_system:

        return path_from_system


    path_from_system = (
        shutil.which("libreoffice")
    )

    if path_from_system:

        return path_from_system


    # Normal Windows locations
    possible_paths = [

        r"C:\Program Files\LibreOffice\program\soffice.exe",

        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",

        r"C:\Program Files\LibreOffice\program\soffice.com",

        r"C:\Program Files (x86)\LibreOffice\program\soffice.com",

    ]


    # Check paths
    for path in possible_paths:

        if os.path.isfile(path):

            return path


    # Try environment variables
    program_files = os.environ.get(
        "PROGRAMFILES"
    )

    program_files_x86 = os.environ.get(
        "PROGRAMFILES(X86)"
    )


    extra_paths = []


    if program_files:

        extra_paths.append(
            os.path.join(
                program_files,
                "LibreOffice",
                "program",
                "soffice.exe"
            )
        )


    if program_files_x86:

        extra_paths.append(
            os.path.join(
                program_files_x86,
                "LibreOffice",
                "program",
                "soffice.exe"
            )
        )


    for path in extra_paths:

        if os.path.isfile(path):

            return path


    return None


# ============================================================
# PPT → PDF
# ============================================================

def convert_ppt_to_pdf(
    input_file,
    output_file
):

    libreoffice = find_libreoffice()


    if not libreoffice:

        raise RuntimeError(

            "LibreOffice was not found. "
            "Please install LibreOffice from "
            "https://www.libreoffice.org/download/"

        )


    temporary_directory = (

        OUTPUT_DIR /
        f"libreoffice_{uuid.uuid4().hex}"

    )


    temporary_directory.mkdir(
        parents=True,
        exist_ok=True
    )


    try:

        command = [

            libreoffice,

            "--headless",

            "--convert-to",
            "pdf",

            "--outdir",
            str(temporary_directory),

            str(input_file)

        ]


        result = subprocess.run(

            command,

            capture_output=True,

            text=True,

            timeout=180

        )


        generated_file = (

            temporary_directory /
            f"{Path(input_file).stem}.pdf"

        )


        if not generated_file.exists():

            error_message = (

                result.stderr.strip()

                or

                result.stdout.strip()

                or

                "LibreOffice failed to convert the PowerPoint file."

            )


            raise RuntimeError(
                error_message
            )


        shutil.copy2(

            generated_file,

            output_file

        )


    finally:

        shutil.rmtree(

            temporary_directory,

            ignore_errors=True

        )


# ============================================================
# IMAGE CONVERSION
# ============================================================

def convert_image(
    input_file,
    output_file,
    output_format
):

    with Image.open(
        input_file
    ) as image:

        if output_format == "JPEG":

            # JPEG does not support transparency
            if image.mode in (
                "RGBA",
                "LA",
                "P"
            ):

                if image.mode == "P":

                    image = image.convert(
                        "RGBA"
                    )


                background = Image.new(
                    "RGB",
                    image.size,
                    "white"
                )


                if "A" in image.getbands():

                    background.paste(

                        image,

                        mask=image.getchannel(
                            "A"
                        )
                    )

                else:

                    background.paste(
                        image
                    )


                image = background

            else:

                image = image.convert(
                    "RGB"
                )


        image.save(

            output_file,

            format=output_format,

            quality=95
        )


# ============================================================
# IMAGE COMPRESSION
# ============================================================

def compress_image(
    input_file,
    output_file,
    quality
):

    with Image.open(
        input_file
    ) as image:

        image_format = image.format


        # ----------------------------------------------------
        # JPEG
        # ----------------------------------------------------

        if image_format in (
            "JPEG",
            "JPG"
        ):

            if image.mode != "RGB":

                image = image.convert(
                    "RGB"
                )


            image.save(

                output_file,

                format="JPEG",

                quality=quality,

                optimize=True,

                progressive=True
            )


        # ----------------------------------------------------
        # PNG
        # ----------------------------------------------------

        elif image_format == "PNG":

            has_transparency = (
                "A" in image.getbands()
            )


            if has_transparency:

                # Preserve transparent PNG
                image.save(

                    output_file,

                    format="PNG",

                    optimize=True,

                    compress_level=9
                )

            else:

                # Non-transparent PNG
                # convert to JPEG for stronger compression
                image = image.convert(
                    "RGB"
                )


                image.save(

                    output_file,

                    format="JPEG",

                    quality=quality,

                    optimize=True,

                    progressive=True
                )


        # ----------------------------------------------------
        # WEBP
        # ----------------------------------------------------

        elif image_format == "WEBP":

            image.save(

                output_file,

                format="WEBP",

                quality=quality,

                method=6
            )


        else:

            raise ValueError(
                "Unsupported image format."
            )


# ============================================================
# PDF → JPG
# ============================================================

def convert_pdf_to_jpg(
    pdf_file,
    output_directory
):

    pdf = fitz.open(
        str(pdf_file)
    )


    pages = []


    for page_number, page in enumerate(
        pdf,
        start=1
    ):

        pixmap = page.get_pixmap(

            matrix=fitz.Matrix(
                1.8,
                1.8
            ),

            alpha=False

        )


        output = (

            output_directory /
            f"page-{page_number}.jpg"

        )


        pixmap.save(
            str(output)
        )


        pages.append(
            output
        )


    pdf.close()


    return pages


# ============================================================
# IMAGES → PDF
# ============================================================

def convert_images_to_pdf(
    image_files,
    output_file
):

    images = []


    try:

        for file in image_files:

            image = Image.open(
                file
            )


            image = image.convert(
                "RGB"
            )


            images.append(
                image
            )


        if not images:

            raise ValueError(
                "No images were selected."
            )


        images[0].save(

            output_file,

            save_all=True,

            append_images=images[1:]

        )


    finally:

        for image in images:

            image.close()


# ============================================================
# MERGE PDF
# ============================================================

def merge_pdf_files(
    pdf_files,
    output_file
):

    merged = fitz.open()


    for file in pdf_files:

        pdf = fitz.open(
            str(file)
        )


        merged.insert_pdf(
            pdf
        )


        pdf.close()


    merged.save(
        str(output_file)
    )


    merged.close()


# ============================================================
# HOME
# ============================================================

@app.route("/")
def home():

    return render_template(
        "index.html"
    )


# ============================================================
# CONVERSION API
# ============================================================

@app.route(
    "/api/convert",
    methods=["POST"]
)
def convert():

    task = request.form.get(
        "task"
    )


    uploaded_files = request.files.getlist(
        "files"
    )


    # --------------------------------------------------------
    # Validate task
    # --------------------------------------------------------

    if task not in ALLOWED:

        return jsonify({

            "error":
            "Invalid conversion tool."

        }), 400


    # --------------------------------------------------------
    # Validate files
    # --------------------------------------------------------

    if (

        not uploaded_files

        or

        not uploaded_files[0].filename

    ):

        return jsonify({

            "error":
            "Please select a file."

        }), 400


    # --------------------------------------------------------
    # Create job directories
    # --------------------------------------------------------

    job_id = uuid.uuid4().hex


    job_upload = (
        UPLOAD_DIR /
        job_id
    )


    job_output = (
        OUTPUT_DIR /
        job_id
    )


    job_upload.mkdir(
        parents=True,
        exist_ok=True
    )


    job_output.mkdir(
        parents=True,
        exist_ok=True
    )


    saved_files = []


    try:

        # ====================================================
        # SAVE UPLOADED FILES
        # ====================================================

        for uploaded in uploaded_files:

            filename = safe_filename(
                uploaded.filename
            )


            extension = (
                Path(filename)
                .suffix
                .lower()
            )


            if extension not in ALLOWED[task]:

                raise ValueError(

                    f"File type {extension} "
                    f"is not supported for this tool."

                )


            path = (
                job_upload /
                filename
            )


            uploaded.save(
                str(path)
            )


            saved_files.append(
                path
            )


        # ====================================================
        # PDF → PPT
        # ====================================================

        if task == "pdf-to-ppt":

            output_file = (

                job_output /
                f"{saved_files[0].stem}.pptx"

            )


            convert_pdf_to_ppt(

                saved_files[0],

                output_file

            )


        # ====================================================
        # PPT → PDF
        # ====================================================

        elif task == "ppt-to-pdf":

            output_file = (

                job_output /
                f"{saved_files[0].stem}.pdf"

            )


            convert_ppt_to_pdf(

                saved_files[0],

                output_file

            )


        # ====================================================
        # JPG → PNG
        # ====================================================

        elif task == "jpg-to-png":

            output_file = (

                job_output /
                f"{saved_files[0].stem}.png"

            )


            convert_image(

                saved_files[0],

                output_file,

                "PNG"

            )


        # ====================================================
        # PNG → JPG
        # ====================================================

        elif task == "png-to-jpg":

            output_file = (

                job_output /
                f"{saved_files[0].stem}.jpg"

            )


            convert_image(

                saved_files[0],

                output_file,

                "JPEG"

            )


        # ====================================================
        # PDF → JPG
        # ====================================================

        elif task == "pdf-to-jpg":

            pages = convert_pdf_to_jpg(

                saved_files[0],

                job_output

            )


            if len(pages) == 1:

                output_file = pages[0]

            else:

                output_file = (

                    job_output /
                    "converted-pages.zip"

                )


                with zipfile.ZipFile(

                    output_file,

                    "w",

                    zipfile.ZIP_DEFLATED

                ) as archive:

                    for page in pages:

                        archive.write(

                            page,

                            page.name

                        )


        # ====================================================
        # JPG → PDF
        # ====================================================

        elif task == "jpg-to-pdf":

            output_file = (

                job_output /
                "converted-images.pdf"

            )


            convert_images_to_pdf(

                saved_files,

                output_file

            )


        # ====================================================
        # PNG → PDF
        # ====================================================

        elif task == "png-to-pdf":

            output_file = (

                job_output /
                "converted-images.pdf"

            )


            convert_images_to_pdf(

                saved_files,

                output_file

            )


        # ====================================================
        # MERGE PDF
        # ====================================================

        elif task == "merge-pdf":

            if len(saved_files) < 2:

                raise ValueError(

                    "Please select at least two PDF files."

                )


            output_file = (

                job_output /
                "merged.pdf"

            )


            merge_pdf_files(

                saved_files,

                output_file

            )


        # ====================================================
        # COMPRESS IMAGE
        # ====================================================

        elif task == "compress-image":

            if len(saved_files) != 1:

                raise ValueError(

                    "Please select exactly one image."

                )


            quality_text = request.form.get(

                "quality",

                "70"

            )


            try:

                quality = int(
                    quality_text
                )

            except (
                ValueError,
                TypeError
            ):

                quality = 70


            quality = max(

                10,

                min(
                    95,
                    quality
                )

            )


            source = saved_files[0]


            extension = (
                source
                .suffix
                .lower()
            )


            # --------------------------------------------
            # Detect transparency
            # --------------------------------------------

            with Image.open(
                source
            ) as image:

                has_alpha = (
                    "A" in image.getbands()
                )


                image_format = image.format


            # --------------------------------------------
            # Output filename
            # --------------------------------------------

            if (

                extension == ".png"

                and

                has_alpha

            ):

                output_file = (

                    job_output /
                    f"{source.stem}-compressed.png"

                )


            elif extension == ".webp":

                output_file = (

                    job_output /
                    f"{source.stem}-compressed.webp"

                )


            else:

                output_file = (

                    job_output /
                    f"{source.stem}-compressed.jpg"

                )


            compress_image(

                source,

                output_file,

                quality

            )


        else:

            raise ValueError(
                "Unknown conversion task."
            )


        # ====================================================
        # VERIFY OUTPUT
        # ====================================================

        if not output_file.exists():

            raise RuntimeError(

                "Conversion completed but "
                "the output file was not created."

            )


        # ====================================================
        # FILE SIZE
        # ====================================================

        original_size = (
            saved_files[0].stat().st_size
        )


        output_size = (
            output_file.stat().st_size
        )


        saved_percent = 0


        if original_size > 0:

            saved_percent = round(

                (

                    1 -

                    (
                        output_size /
                        original_size
                    )

                ) * 100

            )


        # ====================================================
        # RESPONSE
        # ====================================================

        return jsonify({

            "success": True,

            "filename":
                output_file.name,

            "url":
                f"/download/"
                f"{job_id}/"
                f"{output_file.name}",

            "original_size":
                original_size,

            "output_size":
                output_size,

            "saved_percent":
                saved_percent

        })


    except Exception as error:

        print()
        print(
            "========================================"
        )
        print(
            "CONVERSION ERROR"
        )
        print(
            "========================================"
        )
        print(
            str(error)
        )
        print(
            "========================================"
        )
        print()


        return jsonify({

            "error":
            str(error)

        }), 500


    finally:

        # Remove uploaded files
        shutil.rmtree(

            job_upload,

            ignore_errors=True

        )


# ============================================================
# DOWNLOAD
# ============================================================

@app.route(
    "/download/<job_id>/<filename>"
)
def download(
    job_id,
    filename
):

    filename = secure_filename(
        filename
    )


    file_path = (

        OUTPUT_DIR /
        job_id /
        filename

    )


    if not file_path.exists():

        return (
            "File not found.",
            404
        )


    return send_file(

        str(file_path),

        as_attachment=True,

        download_name=file_path.name

    )


# ============================================================
# ERROR HANDLER
# ============================================================

@app.errorhandler(
    413
)
def too_large(error):

    return jsonify({

        "error":
        "The file is too large. "
        "Maximum size is 100 MB."

    }), 413


# ============================================================
# START SERVER
# ============================================================

if __name__ == "__main__":

    print()
    print(
        "=============================================="
    )
    print(
        "          FILEFORGE CONVERTER"
    )
    print(
        "=============================================="
    )
    print()

    libreoffice = find_libreoffice()


    if libreoffice:

        print(
            "LibreOffice found:"
        )

        print(
            libreoffice
        )

    else:

        print(
            "WARNING: LibreOffice not found."
        )

        print(
            "PPT → PDF will not work until LibreOffice is installed."
        )


    print()
    print(
        "Website:"
    )

    print(
        "http://127.0.0.1:5000"
    )

    print()


    app.run(

        host="127.0.0.1",

        port=5000,

        debug=True

    )